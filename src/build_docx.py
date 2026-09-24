"""Convert paper/*.md (simple Markdown) into Value in Health-formatted .docx files.
Formatting per VIH guide: 12-pt Times New Roman, double spacing, 1-inch margins, US Letter,
page numbers bottom-centre, H1 bold, H2 bold italic. Citation markers like [12] or [3,5-7] become superscripts.
Usage: python src/build_docx.py paper/manuscript.md paper/title_page.md ...
"""
import re, sys, pathlib
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CITE = re.compile(r"\[(\d+(?:[-–]\d+)?(?:,\s*\d+(?:[-–]\d+)?)*)\]")
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\^[^^]+\^|~[^~]+~)")

def add_page_number(section):
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; run._r.append(el)

def style_run(run, bold=False, italic=False, sup=False, sub=False, size=12):
    run.font.name = "Times New Roman"; run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.bold = bold; run.italic = italic
    if sup: run.font.superscript = True
    if sub: run.font.subscript = True

def add_rich(par, text, size=12, base_bold=False, base_italic=False):
    # split citation markers first, then inline markup
    pos = 0
    for m in CITE.finditer(text):
        _add_inline(par, text[pos:m.start()], size, base_bold, base_italic)
        style_run(par.add_run(m.group(1).replace(" ", "")), sup=True, size=size)
        pos = m.end()
    _add_inline(par, text[pos:], size, base_bold, base_italic)

def _add_inline(par, text, size, base_bold, base_italic):
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith("**"): style_run(par.add_run(part[2:-2]), bold=True, italic=base_italic, size=size)
        elif part.startswith("*"): style_run(par.add_run(part[1:-1]), italic=True, bold=base_bold, size=size)
        elif part.startswith("^"): style_run(par.add_run(part[1:-1]), sup=True, size=size)
        elif part.startswith("~"): style_run(par.add_run(part[1:-1]), sub=True, size=size)
        else: style_run(par.add_run(part), bold=base_bold, italic=base_italic, size=size)

def build(md_path: pathlib.Path, double=True):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(sec, side, Inches(1))
    add_page_number(sec)
    normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = normal.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE if double else WD_LINE_SPACING.SINGLE

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1; continue
        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|\s*:?-+", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|\s*:?-+", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows)); t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    p = t.cell(r, c).paragraphs[0]; p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    add_rich(p, cell, size=10, base_bold=(r == 0))
            doc.add_paragraph(); continue
        m = re.match(r"^(#{1,3})\s+(.*)", ln)
        if m:
            level = len(m.group(1)); p = doc.add_paragraph()
            if level == 1: style_run(p.add_run(m.group(2)), bold=True, size=12)
            elif level == 2: style_run(p.add_run(m.group(2)), bold=True, italic=True, size=12)
            else: style_run(p.add_run(m.group(2)), italic=True, size=12)
            i += 1; continue
        if re.match(r"^\s*[-*]\s+", ln):
            p = doc.add_paragraph(style="List Bullet"); add_rich(p, re.sub(r"^\s*[-*]\s+", "", ln)); i += 1; continue
        if re.match(r"^\s*\d+\.\s+", ln):
            p = doc.add_paragraph(); add_rich(p, ln.strip()); i += 1; continue
        # paragraph: join consecutive non-empty lines
        buf = [ln]
        while i + 1 < len(lines) and lines[i + 1].strip() and not lines[i + 1].startswith(("#", "|", "- ", "* ")):
            i += 1; buf.append(lines[i].rstrip())
        p = doc.add_paragraph(); add_rich(p, " ".join(buf)); i += 1
    out = md_path.with_suffix(".docx"); doc.save(out); print("wrote", out)

if __name__ == "__main__":
    for a in sys.argv[1:]:
        build(pathlib.Path(a), double=not a.endswith("_tables.md"))
